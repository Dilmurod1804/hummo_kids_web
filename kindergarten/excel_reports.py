import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from decimal import Decimal
from django.utils import timezone
from .models import Child, Group, ChildAttendance, StaffAttendance, MonthlyInvoice, KindergartenSettings

def apply_header_style(cell, text, bg_color="1E1B4B", fg_color="FFFFFF"):
    cell.value = text
    cell.font = Font(name="Segoe UI", size=11, bold=True, color=fg_color)
    cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def apply_title_banner(ws, title_text, max_col, row=1):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
    cell = ws.cell(row=row, column=1)
    cell.value = title_text
    cell.font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="311042", end_color="311042", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 36

def auto_fit_columns(ws, max_col, min_width=14):
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    for col in range(1, max_col + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for cell in ws[col_letter]:
            if cell.row > 1 and cell.value:
                max_len = max(max_len, len(str(cell.value)))
            if cell.row > 2:
                cell.border = thin_border
        ws.column_dimensions[col_letter].width = max(max_len + 4, min_width)


def generate_children_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bolalar Ro'yxati"

    settings = KindergartenSettings.get_settings()
    apply_title_banner(ws, f"{settings.name} — Bolalar va Guruhlar Ro'yxati", 9, row=1)

    # Subtitle / Export Date
    ws.cell(row=2, column=1, value=f"Eksport qilingan sana: {timezone.now().strftime('%d.%m.%Y %H:%M')}")
    ws.cell(row=2, column=1).font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

    headers = [
        "№", "F.I.SH (Bola)", "Guruh", "Yoshi", "Tug'ilgan Sana", 
        "Ota-onasi F.I.SH", "Telefon", "Tibbiy Belgilar / Allergiya", "Holati"
    ]
    
    ws.row_dimensions[3].height = 26
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        apply_header_style(cell, header, "4338CA")

    children = Child.objects.select_related('group').order_by('group__name', 'first_name')
    row_idx = 4
    for idx, child in enumerate(children, 1):
        ws.row_dimensions[row_idx].height = 22
        ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=child.full_name)
        ws.cell(row=row_idx, column=3, value=child.group.name)
        ws.cell(row=row_idx, column=4, value=f"{child.age} yosh").alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=5, value=child.birth_date.strftime("%d.%m.%Y")).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=6, value=child.parent_full_name)
        ws.cell(row=row_idx, column=7, value=child.parent_phone)
        ws.cell(row=row_idx, column=8, value=child.medical_notes or "Belgilar yo'q")
        
        status_cell = ws.cell(row=row_idx, column=9, value="Faol" if child.is_active else "Nofaol")
        status_cell.alignment = Alignment(horizontal="center")
        if child.is_active:
            status_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
            status_cell.font = Font(name="Segoe UI", color="166534", bold=True)
        else:
            status_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            status_cell.font = Font(name="Segoe UI", color="991B1B", bold=True)
        row_idx += 1

    auto_fit_columns(ws, len(headers))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_attendance_excel(group_id=None, date_filter=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Davomat Hisoboti"

    settings = KindergartenSettings.get_settings()
    apply_title_banner(ws, f"{settings.name} — Kunlik Davomat Jurnali", 7, row=1)

    ws.cell(row=2, column=1, value=f"Hisobot sanasi: {date_filter or timezone.now().date()}")
    ws.cell(row=2, column=1).font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

    headers = [
        "№", "Bola F.I.SH", "Guruh", "Sana", "Davomat Holati", "Belgilagan Xodim", "Izoh"
    ]
    ws.row_dimensions[3].height = 26
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        apply_header_style(cell, header, "1E3A8A")

    qs = ChildAttendance.objects.select_related('child', 'child__group', 'marked_by').order_by('-date', 'child__group__name')
    if date_filter:
        qs = qs.filter(date=date_filter)
    if group_id:
        qs = qs.filter(child__group_id=group_id)

    row_idx = 4
    for idx, att in enumerate(qs, 1):
        ws.row_dimensions[row_idx].height = 22
        ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=att.child.full_name)
        ws.cell(row=row_idx, column=3, value=att.child.group.name)
        ws.cell(row=row_idx, column=4, value=att.date.strftime("%d.%m.%Y")).alignment = Alignment(horizontal="center")
        
        status_cell = ws.cell(row=row_idx, column=5, value=att.get_status_display())
        status_cell.alignment = Alignment(horizontal="center")
        if att.status == 'PRESENT':
            status_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
            status_cell.font = Font(name="Segoe UI", color="166534", bold=True)
        elif att.status == 'EXCUSED':
            status_cell.fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
            status_cell.font = Font(name="Segoe UI", color="854D0E", bold=True)
        else:
            status_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            status_cell.font = Font(name="Segoe UI", color="991B1B", bold=True)

        ws.cell(row=row_idx, column=6, value=att.marked_by.get_full_name() if att.marked_by else "Tizim")
        ws.cell(row=row_idx, column=7, value=att.notes or "")
        row_idx += 1

    auto_fit_columns(ws, len(headers))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_staff_attendance_excel(month=None, year=None, staff_id=None, date_val=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Xodimlar Davomati & Jarimalar"

    settings = KindergartenSettings.get_settings()
    apply_title_banner(ws, f"{settings.name} — Xodimlar Davomat, Kechikish va Jarimalar Hisoboti", 12, row=1)

    period_parts = []
    if date_val:
        period_parts.append(f"Sana: {date_val}")
    if month and year:
        period_parts.append(f"Davr: {int(month):02d}/{year}")
    elif year:
        period_parts.append(f"Yil: {year}")
    period_str = " | ".join(period_parts) if period_parts else "Barcha Davrlar"

    ws.cell(row=2, column=1, value=f"Hisobot Davri: {period_str} | Eksport qilingan: {timezone.now().strftime('%d.%m.%Y %H:%M')}")
    ws.cell(row=2, column=1).font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

    headers = [
        "№", "Xodim (F.I.SH)", "Lavozim", "Sana", "Kelgan Vaqti", "Ketgan Vaqti", 
        "Kechikish (daq)", "Barvaqt Ketish (daq)", "Shtraf Summasi (UZS)",
        "Masofa (Bog'chagacha)", "Geofencing Holati", "Tizim Xulosasi"
    ]
    ws.row_dimensions[3].height = 26
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        apply_header_style(cell, header, "0F766E")

    staff_logs = StaffAttendance.objects.select_related('teacher').order_by('-date', '-created_at')
    if date_val:
        staff_logs = staff_logs.filter(date=date_val)
    if month:
        staff_logs = staff_logs.filter(date__month=int(month))
    if year:
        staff_logs = staff_logs.filter(date__year=int(year))
    if staff_id:
        staff_logs = staff_logs.filter(teacher_id=staff_id)

    row_idx = 4
    total_penalties = Decimal('0.00')
    total_late_min = 0

    for idx, log in enumerate(staff_logs, 1):
        ws.row_dimensions[row_idx].height = 22
        ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=log.teacher.get_full_name() or log.teacher.username)
        ws.cell(row=row_idx, column=3, value=log.teacher.position_display).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=4, value=log.date.strftime("%d.%m.%Y")).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=5, value=log.check_in_time.strftime("%H:%M:%S") if log.check_in_time else "—").alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=6, value=log.check_out_time.strftime("%H:%M:%S") if log.check_out_time else "—").alignment = Alignment(horizontal="center")
        
        # Kechikish
        late_cell = ws.cell(row=row_idx, column=7, value=log.late_minutes if log.is_late else 0)
        late_cell.alignment = Alignment(horizontal="center")
        if log.is_late and log.late_minutes > 0:
            late_cell.font = Font(name="Segoe UI", color="DC2626", bold=True)
            total_late_min += log.late_minutes

        # Barvaqt ketish
        early_cell = ws.cell(row=row_idx, column=8, value=log.early_leave_minutes if log.is_early_leave else 0)
        early_cell.alignment = Alignment(horizontal="center")
        if log.is_early_leave and log.early_leave_minutes > 0:
            early_cell.font = Font(name="Segoe UI", color="EA580C", bold=True)

        # Shtraf
        penalty = (log.penalty_amount or Decimal('0.00')) + (log.early_leave_penalty_amount or Decimal('0.00'))
        total_penalties += penalty
        pen_cell = ws.cell(row=row_idx, column=9, value=float(penalty))
        pen_cell.number_format = '#,##0 "UZS"'
        pen_cell.alignment = Alignment(horizontal="right")
        if penalty > 0:
            pen_cell.font = Font(name="Segoe UI", color="DC2626", bold=True)

        ws.cell(row=row_idx, column=10, value=f"{log.distance_meters or 0:.0f} metr").alignment = Alignment(horizontal="center")
        
        geo_cell = ws.cell(row=row_idx, column=11, value="Hudud Ichida" if log.is_within_geofence else "Hududdan Tashqarida")
        geo_cell.alignment = Alignment(horizontal="center")
        if log.is_within_geofence:
            geo_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
            geo_cell.font = Font(name="Segoe UI", color="166534", bold=True)
        else:
            geo_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            geo_cell.font = Font(name="Segoe UI", color="991B1B", bold=True)

        ws.cell(row=row_idx, column=12, value=log.get_status_display()).alignment = Alignment(horizontal="center")
        row_idx += 1

    # Jami qator
    if len(staff_logs) > 0:
        ws.row_dimensions[row_idx].height = 24
        ws.cell(row=row_idx, column=1, value="")
        total_title = ws.cell(row=row_idx, column=2, value="JAMI:")
        total_title.font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
        
        ws.cell(row=row_idx, column=7, value=total_late_min).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=7).font = Font(name="Segoe UI", bold=True, color="DC2626")

        tot_pen = ws.cell(row=row_idx, column=9, value=float(total_penalties))
        tot_pen.number_format = '#,##0 "UZS"'
        tot_pen.font = Font(name="Segoe UI", bold=True, color="DC2626")
        tot_pen.alignment = Alignment(horizontal="right")

    auto_fit_columns(ws, len(headers))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_finance_excel(month=None, year=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Moliya va Qayta Hisoblash"

    settings = KindergartenSettings.get_settings()
    apply_title_banner(ws, f"{settings.name} — Oylik To'lovlar va Qayta Hisoblash (Перерасчет)", 10, row=1)

    period_str = f"{month}/{year}" if month and year else "Barcha Davrlar"
    ws.cell(row=2, column=1, value=f"Hisobot Davri: {period_str} | Kunlik ovqatlanish normasi: {settings.daily_meal_rate:,.0f} UZS")
    ws.cell(row=2, column=1).font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

    headers = [
        "№", "Bola F.I.SH", "Guruh", "Davr (Oy/Yil)", "Asosiy To'lov (UZS)", 
        "Sababli Qoldirilgan (Kun)", "Qayta Hisoblash Chegirmasi (UZS)", 
        "Jami To'lanishi Kerak (UZS)", "Naqd To'lov (UZS)", "Karta / Bank To'lov (UZS)",
        "Jami To'langan (UZS)", "Qarzdorlik / Holat"
    ]
    ws.row_dimensions[3].height = 26
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        apply_header_style(cell, header, "6B21A8")

    invoices = MonthlyInvoice.objects.select_related('child', 'child__group').prefetch_related('payments').order_by('-year', '-month', 'child__first_name')
    if month:
        invoices = invoices.filter(month=month)
    if year:
        invoices = invoices.filter(year=year)

    row_idx = 4
    total_base = 0
    total_recalc = 0
    total_final = 0
    total_cash = 0
    total_card = 0
    total_paid = 0

    for idx, inv in enumerate(invoices, 1):
        inv_cash = float(sum(p.amount for p in inv.payments.all() if p.payment_method == 'CASH'))
        inv_card = float(sum(p.amount for p in inv.payments.all() if p.payment_method in ['CARD', 'BANK_TRANSFER', 'CLICK_PAYME']))

        ws.row_dimensions[row_idx].height = 22
        ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=inv.child.full_name)
        ws.cell(row=row_idx, column=3, value=inv.child.group.name if inv.child.group else "-")
        ws.cell(row=row_idx, column=4, value=f"{inv.month:02d}/{inv.year}").alignment = Alignment(horizontal="center")
        
        c5 = ws.cell(row=row_idx, column=5, value=float(inv.base_fee))
        c5.number_format = '#,##0'
        
        c6 = ws.cell(row=row_idx, column=6, value=inv.excused_days_count)
        c6.alignment = Alignment(horizontal="center")
        
        c7 = ws.cell(row=row_idx, column=7, value=float(inv.recalculation_amount))
        c7.number_format = '#,##0'
        
        c8 = ws.cell(row=row_idx, column=8, value=float(inv.total_amount))
        c8.number_format = '#,##0'
        
        c9 = ws.cell(row=row_idx, column=9, value=inv_cash)
        c9.number_format = '#,##0'

        c10 = ws.cell(row=row_idx, column=10, value=inv_card)
        c10.number_format = '#,##0'

        c11 = ws.cell(row=row_idx, column=11, value=float(inv.paid_amount))
        c11.number_format = '#,##0'

        debt_cell = ws.cell(row=row_idx, column=12)
        debt_cell.alignment = Alignment(horizontal="center")
        if inv.status == 'PAID':
            debt_cell.value = "To'liq To'langan"
            debt_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
            debt_cell.font = Font(name="Segoe UI", color="166534", bold=True)
        elif inv.status == 'PARTIAL':
            debt_cell.value = f"Qisman ({float(inv.remaining_debt):,.0f} UZS qarz)"
            debt_cell.fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
            debt_cell.font = Font(name="Segoe UI", color="854D0E", bold=True)
        else:
            debt_cell.value = f"Qarzdor ({float(inv.remaining_debt):,.0f} UZS)"
            debt_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            debt_cell.font = Font(name="Segoe UI", color="991B1B", bold=True)

        total_base += float(inv.base_fee)
        total_recalc += float(inv.recalculation_amount)
        total_final += float(inv.total_amount)
        total_cash += inv_cash
        total_card += inv_card
        total_paid += float(inv.paid_amount)
        row_idx += 1

    # Total Row
    ws.row_dimensions[row_idx].height = 26
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
    tot_label = ws.cell(row=row_idx, column=1, value="JAMI YIG'INDI:")
    tot_label.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    tot_label.fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    tot_label.alignment = Alignment(horizontal="right", vertical="center")

    for c_idx, val in [(5, total_base), (7, total_recalc), (8, total_final), (9, total_cash), (10, total_card), (11, total_paid)]:
        tc = ws.cell(row=row_idx, column=c_idx, value=val)
        tc.number_format = '#,##0'
        tc.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        tc.fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")

    tot_remain = ws.cell(row=row_idx, column=12, value=f"Qarz: {total_final - total_paid:,.0f} UZS")
    tot_remain.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    tot_remain.fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    tot_remain.alignment = Alignment(horizontal="center")

    auto_fit_columns(ws, len(headers))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_expenses_excel(date_from=None, date_to=None):
    """Chiqimlar (Xarajatlar) Excel hisoboti"""
    from .models import Expense
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chiqimlar"

    settings = KindergartenSettings.get_settings()
    apply_title_banner(ws, f"{settings.name} — Chiqimlar va Xarajatlar Hisoboti", 8, row=1)

    period_str = ""
    if date_from and date_to:
        period_str = f"{date_from} dan {date_to} gacha"
    elif date_from:
        period_str = f"{date_from} dan"
    elif date_to:
        period_str = f"{date_to} gacha"
    ws.cell(row=2, column=1, value=f"Eksport: {settings.timezone if hasattr(settings, 'timezone') else 'O\'zbekiston vaqti'} | Davr: {period_str or 'Barcha davrlar'}")
    ws.cell(row=2, column=1).font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

    headers = [
        "№", "Sana", "Nomi / Sababi", "Kategoriya", "To'lov Usuli",
        "Chek / Hujjat", "Summa (UZS)", "Kiritgan Xodim"
    ]
    ws.row_dimensions[3].height = 26
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        apply_header_style(cell, header, "9F1239")  # Red for expenses

    expenses = Expense.objects.select_related('created_by').order_by('-date')
    if date_from:
        expenses = expenses.filter(date__gte=date_from)
    if date_to:
        expenses = expenses.filter(date__lte=date_to)

    row_idx = 4
    total_amount = 0
    total_cash = 0
    total_card = 0

    for idx, exp in enumerate(expenses, 1):
        ws.row_dimensions[row_idx].height = 22
        ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=exp.date.strftime("%d.%m.%Y")).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=3, value=exp.title)
        ws.cell(row=row_idx, column=4, value=exp.get_category_display()).alignment = Alignment(horizontal="center")

        method_cell = ws.cell(row=row_idx, column=5, value=exp.get_payment_method_display())
        method_cell.alignment = Alignment(horizontal="center")
        if exp.payment_method == 'CASH':
            method_cell.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
            method_cell.font = Font(name="Segoe UI", color="065F46")
        elif exp.payment_method == 'CARD':
            method_cell.fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
            method_cell.font = Font(name="Segoe UI", color="1E40AF")

        ws.cell(row=row_idx, column=6, value=exp.receipt_number or "—").alignment = Alignment(horizontal="center")

        amt_cell = ws.cell(row=row_idx, column=7, value=float(exp.amount))
        amt_cell.number_format = '#,##0'
        amt_cell.font = Font(name="Segoe UI", bold=True, color="9F1239")

        ws.cell(row=row_idx, column=8, value=exp.created_by.get_full_name() if exp.created_by else "—")

        total_amount += float(exp.amount)
        if exp.payment_method == 'CASH':
            total_cash += float(exp.amount)
        elif exp.payment_method == 'CARD':
            total_card += float(exp.amount)
        row_idx += 1

    # Totals row
    ws.row_dimensions[row_idx].height = 26
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
    tot_label = ws.cell(row=row_idx, column=1, value="JAMI CHIQIMLAR:")
    tot_label.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    tot_label.fill = PatternFill(start_color="9F1239", end_color="9F1239", fill_type="solid")
    tot_label.alignment = Alignment(horizontal="right", vertical="center")

    tc = ws.cell(row=row_idx, column=7, value=total_amount)
    tc.number_format = '#,##0'
    tc.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    tc.fill = PatternFill(start_color="9F1239", end_color="9F1239", fill_type="solid")

    auto_fit_columns(ws, len(headers))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

